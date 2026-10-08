from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r'D:\E-waste\E-waste')
OUT = ROOT / 'output' / 'documents' / 'E_Waste_Compute_Grid_Team_Development_Plan.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.65), Inches(.65)
sec.left_margin = sec.right_margin = Inches(.7)
sec.header_distance = sec.footer_distance = Inches(.28)
for name in ['Normal', 'Body Text', 'List Bullet', 'List Number']:
    s = doc.styles[name]
    s.font.name = 'Calibri'; s.font.size = Pt(11)
    s.paragraph_format.space_after = Pt(5)
    s.paragraph_format.line_spacing = 1.05
for name, size in [('Title', 24), ('Heading 1', 16), ('Heading 2', 12.5)]:
    s = doc.styles[name]
    s.font.name = 'Calibri'; s.font.size = Pt(size)
    s.font.color.rgb = RGBColor.from_string('000000' if name == 'Title' else '17365D')
    s.paragraph_format.space_before = Pt(10 if name != 'Title' else 0)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.keep_with_next = True
    # Remove inherited decorative paragraph borders.
    ppr = s.element.find(qn('w:pPr'))
    if ppr is not None:
        for b in list(ppr.findall(qn('w:pBdr'))): ppr.remove(b)
doc.core_properties.title = 'E Waste Compute Grid Team Development Plan'
doc.core_properties.subject = 'Four person ownership and phase delivery plan'
doc.core_properties.author = 'E Waste Compute Grid Team'
header = sec.header.paragraphs[0]
header.add_run('E WASTE COMPUTE GRID  |  TEAM DEVELOPMENT PLAN').font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor.from_string('666666')
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Team working copy  |  Page ').font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def p(text, style=None):
    return doc.add_paragraph(text, style)

def lead(label, text):
    x=doc.add_paragraph(); x.add_run(label + ' ').bold=True; x.add_run(text); return x

def h(text, level=1):
    return doc.add_heading(text, level)

def page():
    doc.add_page_break()

def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    for c,w in zip(t.rows[0].cells,widths): c.width=Inches(w)
    for i,val in enumerate(headers): t.rows[0].cells[i].text=val
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row): cells[i].text=val; cells[i].width=Inches(widths[i])
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+edge); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D9D9D9'); borders.append(x)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge in ['top','left','bottom','right']:
        x=OxmlElement('w:'+edge); x.set(qn('w:w'),'90'); x.set(qn('w:type'),'dxa'); margins.append(x)
    pr.append(margins)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        if ri==0: trpr.append(OxmlElement('w:tblHeader'))
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'17365D' if ri==0 else ('F1F5F9' if ri%2==0 else 'FFFFFF')); cell._tc.get_or_add_tcPr().append(sh)
            for pp in cell.paragraphs:
                pp.paragraph_format.space_after=Pt(2)
                pp.paragraph_format.line_spacing=1.0
                for run in pp.runs:
                    run.font.size=Pt(10)
                    if ri==0: run.bold=True; run.font.color.rgb=RGBColor(255,255,255)
    p('').paragraph_format.space_after=Pt(0)
    return t

def phase(n,title,k,v,ni,c,gate):
    h(f'Phase {n} {title}',2)
    for name,task in [('Kushagra',k),('Vaibhav',v),('Nikhil',ni),('Chinmay',c)]:
        x=lead(name,task); x.paragraph_format.space_after=Pt(3)
    x=lead('Integration gate',gate); x.paragraph_format.space_after=Pt(7)

# Page 1
p('E Waste Compute Grid\nTeam Development Plan','Title')
p('Owners milestones and integration gates')
p('Prepared for Kushagra, Vaibhav, Nikhil and Chinmay\nPlanning baseline 7 October 2026')
h('Our first shared milestone')
p('We will make one real image job work from the dashboard through one physical worker and back to a downloadable result. Then we will expand to three machines, demonstrate worker failure and recovery, and deploy the same verified flow through AWS.')
p('Each member owns a component, but a phase is complete only when the components work together. This document assigns work; it does not claim the proposed changes are already implemented.')
h('Primary ownership')
table(['Member','Owns','Works closely with'],[
 ('Kushagra','Control plane, scheduler, task lifecycle, database and API contracts','Vaibhav on reliability and AWS; Nikhil on scheduling inputs'),
 ('Vaibhav','Worker agent, Docker execution, networking, runtime reliability and deployment','Kushagra on worker protocol; Nikhil on workload runtime'),
 ('Nikhil','AI workloads, genuine datasets, model/runtime validation, benchmarks and energy methodology','Vaibhav on execution; Kushagra on telemetry and scoring'),
 ('Chinmay','React dashboard, API integration, readable results and browser verification','Kushagra on read APIs; all members on demo evidence'),
], [1.0,3.25,2.85])
h('What we can reuse')
p('The existing pipeline includes registration, heartbeats, job splitting, PostgreSQL task claims, inference, uploads, aggregation, leases and retries. The earlier audit passed 38 tests and application builds. Its workers ran locally in native mode; browser behavior, three physical devices, Docker edge execution and AWS deployment remain unverified.')
lead('Architecture decision','Keep PostgreSQL as the queue and use outbound worker polling. Do not add Redis, SQS or a new framework before a demonstrated need.')
lead('Current blockers','Copied Python environment, inconsistent local database settings, synthetic sample photos, missing task ownership/history views and advanced utilities not connected to the live system.')

# Page 2
page(); h('Start together with these four assignments')
p('Begin these preparation tasks together. Integrate as soon as the backend and worker are reachable; do not wait for all components to become polished. There is no fixed deadline until the team confirms device access and availability.')
h('Kushagra establish the integration contract',2)
p('Verify the local database and API topology. Document the existing upload, job, claim, heartbeat and completion bodies, including authentication, idempotency keys and lease tokens. Strengthen result validation to account for the exact assigned image identities.')
lead('Files','control-plane/src/config.ts, app.ts, validation.ts; prisma/schema.prisma; docs/api_contracts.md; control-plane/tests.')
lead('Deliver','A reachable backend, tested request examples and a result contract reviewed by Vaibhav, Nikhil and Chinmay.')
h('Vaibhav prepare the first real worker',2)
p('Create a fresh supported Python environment. Build or load both approved workload images, verify Docker and paths on one physical machine, and configure its reachable backend address. Keep the worker identity private and distinct.')
lead('Files','worker/config.py, api_client.py, task_executor.py, README.md, .env.example; infra/docker-compose.yml and Dockerfiles.')
lead('Deliver','A worker that registers, verifies its real runtime, downloads inputs, executes and returns results without manual task operation.')
h('Nikhil replace placeholders with useful inputs',2)
p('Prepare ten genuine photographs with expected labels and provenance. Make download failure visible instead of silently saving blank images. Verify model/labels, record real predictions, and explain errors or low confidence honestly.')
lead('Files','workloads/test-data/setup_sample_data.py and sample data; ai-inference/run_inference.py; model/runtime documentation.')
lead('Deliver','A genuine input manifest and valid result examples handed to Kushagra and Chinmay; runtime instructions handed to Vaibhav.')
h('Chinmay verify the user journey',2)
p('Connect the dashboard to the real backend. Clearly distinguish the pipeline demo from image classification. Verify login, upload, submission, live progress and result download; show meaningful connection and submission errors.')
lead('Files','frontend/src/App.tsx, services/api.ts, App.css; frontend/vite.config.ts; browser verification notes.')
lead('Deliver','A repeatable browser flow using Nikhil’s photos and Vaibhav’s worker, with screenshots and sanitized job IDs.')
lead('Shared acceptance','Ten real images, chunk size five, two tasks, one physical worker, every input counted once as a prediction or explicit error, and a downloaded result. Relevant checks must run with database integration enabled.')

# Page 3
page(); h('Phases 0 to 3 establish the working slice')
p('These phase numbers match the master project plan. Keep working registration and queue code; fix demonstrated defects and add the missing evidence.')
phase(0,'Recover the reproducible baseline',
 'Recover the canonical Git checkout; choose database topology; document backend startup and environment variables.',
 'Recreate worker setup; verify Docker/images, identity paths, power, networking and sleep settings.',
 'Verify model/labels/dependencies and separate real demonstration data from synthetic fixtures.',
 'Verify frontend build and API proxy; document how a teammate connects.',
 'A teammate can start the backend and one registered worker from clean instructions without the copied virtual environment.')
phase(1,'Validate real inputs and contracts',
 'Own API examples and exact assigned-input result validation; add meaningful backend regression tests.',
 'Verify workload capability before advertising it; align worker request and reporting bodies with the contract.',
 'Curate labeled real photos; remove silent realistic-looking fallbacks; validate classifier output and schema.',
 'Make workload modes clear; align upload/job/result types and user-visible validation errors.',
 'Real inputs produce valid predictions or explicit errors; clients share the same reviewed contract.')
phase(2,'Complete the first browser and worker flow',
 'Trace upload, creation, claim, acceptance and aggregation; retain task/job IDs and check counters.',
 'Run the actual approved container on a physical device; verify staging, upload, cleanup and heartbeat.',
 'Confirm the actual model output on the submitted ten-photo dataset and distinguish accuracy from delivery.',
 'Drive and record dashboard submission, progress and result download against the real backend.',
 'The ten-photo two-task acceptance run passes and another teammate can repeat it.')
phase(3,'Add task and assignment history',
 'Design and migrate TaskAttempt history and timing; expose bounded authenticated job/task/attempt read APIs.',
 'Report agreed phase/timing metadata and verify lease/attempt handling across retries and reconnects.',
 'Validate original image identity and timing/metric meaning; provide readable result examples.',
 'Build job detail, task-owner/retry timeline and prediction table using the new APIs.',
 'Who ran each task and what happened to every attempt is visible and survives a control-plane restart.')
lead('Review responsibility','Kushagra leads phase integration; Vaibhav reviews worker-facing APIs, Chinmay reviews dashboard-facing APIs, and Nikhil reviews workload/result semantics. Schema changes require migration and concurrency review.')

# Page 4
page(); h('Phases 4 to 8 prove the hackathon system')
phase(4,'Run three machines and recover a failure',
 'Verify compatible concurrent claims, expired-lease requeue and rejection of late results.',
 'Prepare three distinct physical identities; abruptly kill an assigned worker and verify restart cleanup.',
 'Choose enough genuine work to show overlap; record hardware and runtime conditions.',
 'Show the lost worker, affected task, new owner/attempt and final result.',
 'Three physical machines contribute; unfinished work is recovered elsewhere with one accepted result.')
phase(5,'Report performance and energy honestly',
 'Store agreed attempt timings/provenance; separate job elapsed time from summed task durations.',
 'Connect Nikhil’s measurement adapter to execution; collect agreed runtime metadata.',
 'Own baseline calibration, optional actual meter, throughput experiment and measurement methodology.',
 'Display job/node contribution, throughput and energy with units, scope and estimated/measured labels.',
 'Numbers have reproducible methods; excluded costs and uncertainties are stated; no assumed cloud savings.')
phase(6,'Protect the service before public exposure',
 'Own authorization, result validation, credential lifecycle, limits and safe API errors.',
 'Own worker credential storage, runtime isolation, endpoint checks and TLS/network setup.',
 'Verify workload/model provenance and safe input decoding/resource requirements.',
 'Handle unauthorized/degraded states; keep secrets out of frontend builds and shared evidence.',
 'Public deployment uses unique credentials, HTTPS, scoped access and private data services.')
phase(7,'Make failures diagnosable',
 'Add correlated structured logs, readiness/liveness, backlog and recovery-sweep metrics.',
 'Collect sanitized worker/container logs and configure retention and useful alerts.',
 'Validate throughput/energy definitions and identify measurement coverage limits.',
 'Expose fresh/stale state and understandable errors; verify failure information is visible.',
 'A teammate can diagnose no-worker, corrupt-input and storage-failure stalls from retained evidence.')
phase(8,'Deploy and repeat through AWS',
 'Own backend/database migration and API/storage/IAM architecture; review deployment security.',
 'Lead deployment scripts, HTTPS, networking, secret injection, logs, rollback and teardown.',
 'Validate the model/runtime on cloud-connected physical workers and real S3 transfers.',
 'Serve the built dashboard properly and verify the complete deployed browser flow.',
 'The same real workload and recovery demonstration passes through deployed HTTPS and actual S3.')

# Page 5
page(); h('Phases 9 to 12 strengthen the portfolio')
p('The hackathon release can freeze after phases 0 to 8 pass. Do not wait for every advanced feature. Portfolio work builds on the verified release without destabilizing it.')
phase(9,'Make scheduling explainable',
 'Implement fresh load/benchmark eligibility and a live scheduling policy; persist inputs and rationale.',
 'Report bounded fresh CPU/RAM/runtime telemetry and enforce local refusal when capacity is insufficient.',
 'Measure per-workload throughput under real resource limits; validate policy inputs and comparisons.',
 'Show actual eligibility and selection rationale, with measurement timestamp and policy version.',
 'The policy affects live decisions and is compared with FIFO; fixture scores are not presented as live scheduling.')
phase(10,'Improve operations and maintainability',
 'Extract backend responsibilities; own reviewed migrations, constraints, admission and optional cancellation.',
 'Add service install/start/stop, drain/reconnect and cleanup; rehearse backup/restore and deployments.',
 'Verify runtime compatibility and repeatable measurement across supported devices.',
 'Test management controls, pagination, keyboard/accessibility and browser error paths.',
 'Clean setup, relevant CI, Docker/browser/fault tests, five-worker soak and restore/rollback evidence pass.')
phase(11,'Add a second workload and optional energy policy',
 'Define versioned splitter/validator/aggregator adapters; preserve common lease/retry behavior.',
 'Integrate the approved workload runtime and capability advertisement without arbitrary code execution.',
 'Build a genuine second independent workload; supply trusted benchmark/energy inputs for any optimizer.',
 'Add workload-specific input/results views and explain any enabled policy.',
 'Two real workloads share orchestration; energy optimization is claimed only when measured and used live.')
phase(12,'Freeze and publish the verified release',
 'Freeze contracts and tag the canonical release; finalize architecture, decisions, API and limits.',
 'Finalize preflight/startup/deployment instructions and rehearse actual worker failure.',
 'Prepare reproducible dataset/benchmark/energy evidence and the hardware reuse explanation.',
 'Lead the 3–5 minute dashboard demonstration, screenshots and recorded fallback.',
 'Demo and portfolio claims match evidence; no untested feature is presented as complete.')
h('Work we deliberately defer',2)
p('Redis/SQS, global worker reservations, adaptive chunking, live push updates and public volunteer execution require a demonstrated need and separate acceptance criteria. Image classification is the first supported real workload; a second workload should follow a stable core release.')

# Page 6
page(); h('Handoffs and shared interfaces')
p('Agree on these deliverables before dependent implementation. Share a validated example and executed check, not only a screenshot or verbal description.')
table(['From','To','Handoff'],[
 ('Kushagra','Vaibhav','Registration, heartbeat, claim, upload and complete/fail bodies; auth; lease/retry semantics; test API URL.'),
 ('Kushagra','Chinmay','Typed job/node/task/attempt/results responses; pagination, error cases and metric definitions.'),
 ('Nikhil','Vaibhav','Approved image/tag or digest; command; input/output paths; model/labels; resource requirements and runtime preflight.'),
 ('Nikhil','Kushagra','Input identity/result schema; genuine data; versioned throughput and energy source/scope.'),
 ('Nikhil','Chinmay','Expected prediction/error examples, interpretation of confidence and measurement labels.'),
 ('Vaibhav','All','Clean setup/deploy commands; reachable addresses; device inventory, release version and sanitized runtime evidence.'),
 ('Chinmay','All','Reproducible browser/demo script, screenshots and failures found at the API integration boundary.'),
], [1.0,1.0,5.1])
h('File ownership prevents conflicting changes',2)
p('Kushagra owns control-plane source, Prisma and shared API contracts. Vaibhav owns worker orchestration/execution and infra. Nikhil owns workload scripts, datasets, benchmarks and energy utilities. Chinmay owns frontend source and browser tests. Work on another owner’s file through a reviewed change; coordinate before simultaneous edits.')
lead('Important boundary','Nikhil owns measurement logic; Vaibhav connects it to worker execution; Kushagra stores and aggregates it; Chinmay displays it. No member independently invents a different energy schema.')
h('Contract changes and review',2)
p('Document changed requests/responses, statuses, units and version impact before merging. The primary owner proposes the change; every affected client owner reviews it. Keep compatibility where reasonable and add a regression test for any behavior change.')
lead('Required reviewers','Worker protocol: Kushagra and Vaibhav. Workload output: Nikhil and Kushagra. Dashboard reads: Kushagra and Chinmay. Runtime packaging: Vaibhav and Nikhil. Deployment/security: Kushagra and Vaibhav.')
lead('Data rules','Never share tokens, presigned URLs, credential directories or database dumps in team updates. Synthetic results are labeled. Unknown energy is not zero energy. Repeated execution may happen; accepted results must not be counted twice.')

# Page 7
page(); h('Integration workflow and release checklist')
h('Daily working rhythm',2)
p('Start with a short check-in: what changed, what is blocked, and which shared contract is affected. Work in small reviewable branches in the canonical repository. Integrate toward the current slice as soon as a component is usable. End with a short live run or executed check and an update to the evidence ledger.')
lead('Status update format','Member; task/phase; changed files; contract/schema impact; checks executed and environment; result; blocker; next step. Label unexecuted checks explicitly.')
lead('Done means','Implementation, relevant validation, an integration demonstration, current documentation and review by an affected owner. A file existing or a test using mocks is not sufficient evidence of a real feature.')
h('Who verifies what',2)
table(['Owner','Primary checks','Integration partner'],[
 ('Kushagra','API/database/scheduler concurrency, idempotency, leases, exact result validation and migrations','Vaibhav for protocol; Nikhil for result semantics'),
 ('Vaibhav','Real Docker execution, paths/limits, reconnect, process kill, cleanup, networking and deployment','Kushagra for recovery; Nikhil for runtime'),
 ('Nikhil','Genuine data, inference correctness, corrupt-input handling, benchmark and energy methods','Vaibhav for actual execution'),
 ('Chinmay','Frontend build/lint, browser workflow, progress/results, errors, accessibility and demo','Kushagra for API consistency'),
], [1.0,4.0,2.1])
p('Backend integration requires TEST_DATABASE_URL and a working TEST_PYTHON when the default environment is unavailable. Use a disposable database. Record the actual suite and platform; a successful command can still have skipped database integration.')
h('Hackathon release gate',2)
for item in [
 'Three distinct authorized physical machines contribute to one genuine job with overlapping execution.',
 'An assigned worker is abruptly killed; another machine finishes its task under a new attempt.',
 'Every input is accounted for exactly once in accepted results; failures and partial success are explicit.',
 'Dashboard shows task ownership/retry and readable output, with correct timing and energy labels.',
 'Clean setup and relevant checks pass; AWS is demonstrated only if real deployed transfers were verified.',
 'A rehearsed 3–5 minute demonstration and clearly labeled fallback match the frozen release.',
]: p(item,'List Bullet')
h('Portfolio follow through',2)
p('Retain reviewed code, migrations, CI/browser/Docker/fault evidence, restore/rollback instructions, architecture/API documentation, same-work benchmark methods, screenshots/video and known limits. Keep the master plan evidence ledger current after every completed phase.')
lead('Next team checkpoint','Run the ten-photo two-task browser-to-physical-worker slice together. Kushagra coordinates the gate; Vaibhav operates the worker, Nikhil verifies outputs and Chinmay drives the dashboard.')

doc.save(OUT)
print(OUT)
print('Paragraphs:',len(doc.paragraphs),'Tables:',len(doc.tables))
